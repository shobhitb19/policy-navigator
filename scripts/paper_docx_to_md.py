#!/usr/bin/env python3
"""Convert the position paper .docx to a parallel .md for reading on the repo.

Usage: python3 scripts/paper_docx_to_md.py "outputs/Cities and Spatial Policy 08.docx"

Writes a .md file alongside the .docx with the same basename. Word's TOC is
skipped (GitHub renders headings); tables become Markdown tables; figure
images are replaced with an italic placeholder line under their caption.
"""

import sys
import os
from docx import Document
from docx.oxml.ns import qn

HEADING_MAP = {
    "Heading 1": "#",
    "Heading 2": "##",
    "Heading 3": "###",
    "Heading 4": "###",
    "Title": "#",
}

SKIP_STYLES = {"toc 1", "toc 2", "toc 3", "TOC Heading"}


def para_text(p):
    return p.text.strip()


def para_md_text(p):
    """Paragraph text with run-level bold/italic rendered as md markers."""
    out = []
    for r in p.runs:
        t = r.text
        if not t:
            continue
        if r.bold and r.italic:
            out.append(("bi", t))
        elif r.bold:
            out.append(("b", t))
        elif r.italic:
            out.append(("i", t))
        else:
            out.append(("n", t))
    # merge adjacent runs of same format
    merged = []
    for fmt, t in out:
        if merged and merged[-1][0] == fmt:
            merged[-1] = (fmt, merged[-1][1] + t)
        else:
            merged.append((fmt, t))
    marks = {"n": ("", ""), "b": ("**", "**"), "i": ("_", "_"), "bi": ("**_", "_**")}
    parts = []
    for fmt, t in merged:
        lead = t[: len(t) - len(t.lstrip())]
        trail = t[len(t.rstrip()):]
        core = t.strip()
        if not core:
            parts.append(t)
            continue
        a, b = marks[fmt]
        parts.append(f"{lead}{a}{core}{b}{trail}")
    return "".join(parts).strip()


def table_to_md(tbl):
    rows = []
    for r in tbl.rows:
        cells = [c.text.strip().replace("\n", " ").replace("|", "\\|") for c in r.cells]
        rows.append(cells)
    if not rows:
        return ""
    # Single-cell tables are call-out boxes, not data tables
    if len(rows) == 1 and len(rows[0]) == 1:
        body = rows[0][0]
        return "\n".join("> " + line for line in body.split(". ")) + "\n"
    width = max(len(r) for r in rows)
    out = []
    header, *rest = rows
    header += [""] * (width - len(header))
    out.append("| " + " | ".join(header) + " |")
    out.append("|" + "---|" * width)
    for r in rest:
        r += [""] * (width - len(r))
        if not any(c for c in r):
            continue
        out.append("| " + " | ".join(r) + " |")
    return "\n".join(out) + "\n"


def convert(docx_path):
    doc = Document(docx_path)
    body = doc.element.body
    tables = iter(doc.tables)
    paras = iter(doc.paragraphs)

    lines = []
    skipping_contents = False
    seen_heading = False

    for child in body:
        tag = child.tag.split("}")[-1]
        if tag == "tbl":
            lines.append(table_to_md(next(tables)))
            continue
        if tag != "p":
            continue
        p = next(paras)
        style = p.style.name
        text = para_text(p)

        if style in SKIP_STYLES:
            continue
        if text == "Contents":
            skipping_contents = True
            continue
        if skipping_contents:
            # contents block ends at the first real heading
            if style in HEADING_MAP and text:
                skipping_contents = False
            else:
                continue

        has_image = child.find(".//" + qn("w:drawing")) is not None

        if not text:
            # skip decorative images before the body starts (e.g. header logo)
            if has_image and seen_heading:
                lines.append("*[Figure: see Word version for image]*\n")
            continue

        if style in HEADING_MAP:
            seen_heading = True
            lines.append(f"{HEADING_MAP[style]} {text}\n")
        elif style == "Caption":
            lines.append(f"***{text}***\n")
        elif style == "List Paragraph":
            lines.append(f"- {para_md_text(p) or text}\n")
        else:
            lines.append(f"{para_md_text(p) or text}\n")

    return "\n".join(lines).rstrip() + "\n"


def main():
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(1)
    docx_path = sys.argv[1]
    md_path = os.path.splitext(docx_path)[0] + ".md"
    md = convert(docx_path)
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(md)
    print(f"Written: {md_path}")


if __name__ == "__main__":
    main()
